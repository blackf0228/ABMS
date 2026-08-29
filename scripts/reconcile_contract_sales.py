from __future__ import annotations

import subprocess
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(r"C:\Users\black\Documents\Hello World\abms")
OUTPUT = ROOT / "output" / "契約売上突合差異.xlsx"


def run_sql(sql: str) -> list[list[str]]:
    result = subprocess.run(
        [
            "docker",
            "exec",
            "abms-postgres",
            "psql",
            "-U",
            "abms",
            "-d",
            "abms",
            "-At",
            "-F",
            "\t",
            "-c",
            sql,
        ],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    return [line.split("\t") for line in result.stdout.splitlines() if line]


def dec(value: str) -> Decimal:
    return Decimal(value or "0")


def month_start(value: date) -> date:
    return value.replace(day=1)


def next_month(value: date) -> date:
    if value.month == 12:
        return date(value.year + 1, 1, 1)
    return date(value.year, value.month + 1, 1)


def active(contract: dict, month: date) -> bool:
    if contract["start"] is None:
        return False
    return (
        contract["start"] <= month
        and (contract["end"] is None or contract["end"] >= month)
    )


def parse_date(value: str) -> date | None:
    return date.fromisoformat(value) if value else None


def main() -> None:
    contracts = []
    for row in run_sql(
        """
        select c.contract_no, c.engineer_name,
               coalesce(nullif(trim(concat_ws(' ', e.last_name, e.first_name)), ''), e.employee_name, c.engineer_name),
               c.contract_amount, c.engineer_unit_price,
               c.customer_name, coalesce(customer.company_id, ''),
               c.engineer_company_name, coalesce(engineer_company.company_id, ''),
               c.contract_start_date, c.contract_end_date
        from contracts c
        left join employee_masters e on e.employee_no = c.engineer_name
        left join customer_masters customer on customer.company_name = c.customer_name
        left join customer_masters engineer_company on engineer_company.company_name = c.engineer_company_name
        order by c.contract_no
        """
    ):
        contracts.append(
            {
                "contract_no": row[0],
                "engineer_no": row[1],
                "engineer_name": row[2],
                "contract_amount": dec(row[3]),
                "engineer_unit_price": dec(row[4]),
                "customer_name": row[5] or "",
                "customer_id": row[6],
                "engineer_company_name": row[7] or "",
                "engineer_company_id": row[8],
                "start": parse_date(row[9]),
                "end": parse_date(row[10]),
            }
        )

    sales = []
    for row in run_sql(
        """
        select s.no, s.billing_month, s.billing_source_company_id,
               s.billing_destination_company_id, s.engineer_name,
               coalesce(nullif(trim(concat_ws(' ', e.last_name, e.first_name)), ''), e.employee_name, s.engineer_name),
               s.billing_amount,
               coalesce(src.company_name, s.billing_source_company_id),
               coalesce(dst.company_name, s.billing_destination_company_id)
        from sales_imports s
        left join employee_masters e on e.employee_no = s.engineer_name
        left join customer_masters src on src.company_id = s.billing_source_company_id
        left join customer_masters dst on dst.company_id = s.billing_destination_company_id
        order by s.billing_month, s.no
        """
    ):
        billing_month = parse_date(row[1])
        sales.append(
            {
                "no": row[0],
                "month": month_start(billing_month),
                "source_id": row[2],
                "destination_id": row[3],
                "engineer_no": row[4],
                "engineer_name": row[5],
                "amount": dec(row[6]),
                "source_name": row[7],
                "destination_name": row[8],
            }
        )

    if not sales:
        raise RuntimeError("売上データがありません。")

    sales_min = min(item["month"] for item in sales)
    sales_max = max(item["month"] for item in sales)
    sales_diffs = []
    sales_by_key = defaultdict(list)

    for item in sales:
        is_revenue = item["source_id"] == "0000"
        sales_type = "売上" if is_revenue else "BP支払"
        company_id = item["destination_id"] if is_revenue else item["source_id"]
        candidates = [
            contract
            for contract in contracts
            if contract["engineer_no"] == item["engineer_no"]
            and (
                (is_revenue and contract["customer_id"] == company_id)
                or (
                    not is_revenue
                    and contract["engineer_company_id"] == company_id
                )
            )
            and active(contract, item["month"])
        ]
        expected_values = [
            contract["contract_amount"] if is_revenue else contract["engineer_unit_price"]
            for contract in candidates
        ]
        if not candidates:
            reason = "契約なしまたは契約期間外"
        elif len(candidates) > 1:
            reason = "複数契約候補"
        elif item["amount"] != expected_values[0]:
            reason = "金額差"
        else:
            sales_by_key[(item["engineer_no"], item["month"], sales_type, company_id)].append(item)
            continue

        sales_diffs.append(
            [
                reason,
                sales_type,
                item["no"],
                item["month"].strftime("%Y/%m"),
                item["engineer_no"],
                item["engineer_name"],
                item["source_id"],
                item["source_name"],
                item["destination_id"],
                item["destination_name"],
                item["amount"],
                expected_values[0] if len(expected_values) == 1 else "",
                ",".join(contract["contract_no"] for contract in candidates),
            ]
        )

    contract_diffs = []
    for contract in contracts:
        if contract["start"] is None:
            continue
        current = month_start(max(contract["start"], sales_min))
        last = month_start(min(contract["end"] or sales_max, sales_max))
        if current > last:
            continue
        while current <= last:
            streams = [
                ("売上", contract["customer_id"], contract["customer_name"], contract["contract_amount"]),
            ]
            if (
                contract["engineer_company_id"]
                and contract["engineer_company_name"] != "株式会社CMCソフト"
            ):
                streams.append(
                    (
                        "BP支払",
                        contract["engineer_company_id"],
                        contract["engineer_company_name"],
                        contract["engineer_unit_price"],
                    )
                )
            for sales_type, company_id, company_name, expected in streams:
                actual_rows = [
                    item
                    for item in sales
                    if item["engineer_no"] == contract["engineer_no"]
                    and item["month"] == current
                    and (
                        (sales_type == "売上" and item["source_id"] == "0000" and item["destination_id"] == company_id)
                        or (
                            sales_type == "BP支払"
                            and item["source_id"] == company_id
                            and item["destination_id"] == "0000"
                        )
                    )
                ]
                actual = sum((item["amount"] for item in actual_rows), Decimal("0"))
                if not actual_rows:
                    reason = "売上未計上"
                elif actual != expected:
                    reason = "金額差"
                else:
                    current_stream_key = (contract["engineer_no"], current, sales_type, company_id)
                    if current_stream_key not in sales_by_key:
                        sales_by_key[current_stream_key] = actual_rows
                    continue
                contract_diffs.append(
                    [
                        reason,
                        sales_type,
                        contract["contract_no"],
                        current.strftime("%Y/%m"),
                        contract["engineer_no"],
                        contract["engineer_name"],
                        company_id,
                        company_name,
                        expected,
                        actual if actual_rows else "",
                        contract["customer_name"],
                        contract["engineer_company_name"],
                    ]
                )
            current = next_month(current)

    summary = [
        ["項目", "件数"],
        ["契約件数", len(contracts)],
        ["売上明細件数", len(sales)],
        ["売上側差異件数", len(sales_diffs)],
        ["契約側差異件数", len(contract_diffs)],
        ["売上側差異合計金額", sum((row[10] for row in sales_diffs if isinstance(row[10], Decimal)), Decimal("0"))],
    ]

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "サマリー"
    for row in summary:
        summary_sheet.append(row)

    sales_sheet = workbook.create_sheet("売上側差異")
    sales_sheet.append(
        [
            "差異区分", "売上区分", "売上No", "請求年月", "社員番号", "技術者名",
            "請求元会社ID", "請求元会社名", "請求先会社ID", "請求先会社名",
            "売上金額", "契約基準金額", "契約番号候補",
        ]
    )
    for row in sales_diffs:
        sales_sheet.append(row)

    contract_sheet = workbook.create_sheet("契約側差異")
    contract_sheet.append(
        [
            "差異区分", "売上区分", "契約番号", "対象年月", "社員番号", "技術者名",
            "対象会社ID", "対象会社名", "契約基準金額", "売上実績金額",
            "取引先", "技術者所属会社",
        ]
    )
    for row in contract_diffs:
        contract_sheet.append(row)

    header_fill = PatternFill("solid", fgColor="FFF2CC")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 30)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"contracts={len(contracts)} sales={len(sales)} sales_diffs={len(sales_diffs)} contract_diffs={len(contract_diffs)}")


if __name__ == "__main__":
    main()
